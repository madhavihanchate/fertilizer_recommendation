import os, re, openpyxl
from docx import Document
from django.conf import settings

DATA_DIR = os.path.join(settings.BASE_DIR, 'data')
RDF_RULE_BOOK = os.path.join(DATA_DIR, 'RDF_Rule_Book.docx')
CROP_DATASET = os.path.join(DATA_DIR, 'India_State_Wise_Crops_with_Soil_Parameters.xlsx')
FERT_DATASET = os.path.join(DATA_DIR, 'India_Fertilizer_Recommendation_Threshold_Based.xlsx')

def load_rules():
    doc = Document(RDF_RULE_BOOK)
    rules = {}
    param_names = [
        'ph', 'ec', 'organic_carbon', 'nitrogen', 'phosphorus',
        'potassium', 'sulphur', 'calcium', 'magnesium', 'zinc',
        'iron', 'manganese', 'copper', 'boron'
    ]
    for ti, table in enumerate(doc.tables):
        if ti >= len(param_names):
            break
        param = param_names[ti]
        rules[param] = []
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 3:
                rules[param].append({
                    'status': cells[0],
                    'practice': cells[1],
                    'dose': cells[2],
                })
    return rules

def load_crop_dataset():
    wb = openpyxl.load_workbook(CROP_DATASET, data_only=True)
    states = [s for s in wb.sheetnames if s not in ('Soil_Methodology', 'INDEX')]
    data = {}
    for state in states:
        ws = wb[state]
        records = []
        for row in ws.iter_rows(min_row=6, values_only=True):
            vals = [str(v).strip() if v is not None else '' for v in row]
            if not vals[1] or vals[0] == 'Category':
                continue
            records.append({
                'category': vals[0],
                'crop': vals[1],
                'agro_zone': vals[3],
                'oc': _float(vals[4]),
                'n': _float(vals[5]),
                'p2o5': _float(vals[6]),
                'k': _float(vals[7]),
                's': _float(vals[8]),
                'fe': _float(vals[9]),
                'mn': _float(vals[10]),
                'cu': _float(vals[11]),
                'zn': _float(vals[12]),
                'b': _float(vals[13]),
                'ph': _float(vals[14]),
                'ec': _float(vals[15]),
                'ca': _float(vals[16]),
                'mg': _float(vals[17]),
                'rdf_ratio': vals[18],
            })
        data[state] = records
    wb.close()
    return data

def _parse_npk(ratio_str):
    parts = ratio_str.replace(':', ' ').replace('/', ' ').split()
    nums = []
    for p in parts:
        try:
            nums.append(float(p))
        except ValueError:
            continue
    while len(nums) < 3:
        nums.append(0.0)
    return nums[0], nums[1], nums[2]

def _calc_urea(n):
    return round(n / 0.46, 2)

def _calc_dap(p2o5):
    return round(p2o5 / 0.46, 2)

def _calc_mop(k2o):
    return round(k2o / 0.60, 2)

def load_fertilizer_dataset():
    wb = openpyxl.load_workbook(FERT_DATASET, data_only=True)
    states = [s for s in wb.sheetnames if s not in ('Threshold Reference', 'Soil_Methodology', 'INDEX')]
    data = {}
    for state in states:
        ws = wb[state]
        records = []
        for row in ws.iter_rows(min_row=8, values_only=True):
            vals = [str(v).strip() if v is not None else '' for v in row]
            if not vals[1] or vals[0] == 'Category':
                continue
            n, p, k = _parse_npk(vals[5])
            # Read actual values from file columns (9=DAP, 11=Urea, 12=MOP)
            # Fall back to derivation from NPK ratio when actual value is 0
            dap_actual = _float(vals[9])
            urea_actual = _float(vals[11])
            mop_actual = _float(vals[12])
            records.append({
                'category': vals[0],
                'crop': vals[1],
                'state': vals[3],
                'agro_zone': vals[4],
                'npk_ratio': vals[5],
                'urea': urea_actual if urea_actual else _calc_urea(n),
                'dap': dap_actual if dap_actual else _calc_dap(p),
                'mop': mop_actual if mop_actual else _calc_mop(k),
            })
        data[state] = records
    wb.close()
    return data

def _float(v):
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0
